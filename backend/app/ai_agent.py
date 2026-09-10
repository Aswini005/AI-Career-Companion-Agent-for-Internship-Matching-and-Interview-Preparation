import re
from pathlib import Path
import json
import fitz  # PyMuPDF
import docx
from app.config import settings


def extract_doc_text(file_path: str, file_ext: str) -> str:
    path = Path(file_path)
    if not path.exists():
        raise FileNotFoundError(f"File not found: {file_path}")
    ext = file_ext.lower().lstrip(".")
    text = ""
    if ext == "pdf":
        doc = fitz.open(str(path))
        for page in doc:
            text += page.get_text() + "\n"
    elif ext == "docx":
        doc = docx.Document(str(path))
        for p in doc.paragraphs:
            if p.text:
                text += p.text + "\n"
        for table in doc.tables:
            for row in table.rows:
                row_txt = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
                if row_txt:
                    text += row_txt + "\n"
    else:
        raise ValueError(f"Unsupported file format .{ext}")
    return text.strip()

def recommend_roles(resume_data: dict) -> dict:
    skills = [s.strip() for s in resume_data.get("skills", []) if s.strip()]
    tech_skills = [s.strip() for s in resume_data.get("technical_skills", []) if s.strip()]
    all_skills = list(set(skills + tech_skills))
    skills_lower = [s.lower() for s in all_skills]
    
    roles_catalog = [
        {
            "role": "AI / Machine Learning Engineer",
            "required_keywords": ["python", "machine learning", "pytorch", "tensorflow", "fastapi", "pandas", "scikit-learn", "data science", "rag"],
            "suitable_internships": ["AI/ML Intern", "Data Science Intern", "NLP Engineering Intern"],
            "summary": "Building AI models, LLM pipelines, and data analytics systems."
        },
        {
            "role": "Full Stack Software Developer",
            "required_keywords": ["react", "javascript", "html", "css", "node", "fastapi", "python", "sql", "rest api", "git"],
            "suitable_internships": ["Software Engineering Intern", "Full Stack Intern", "Frontend Intern"],
            "summary": "Designing modern Web UIs and robust REST APIs."
        },
        {
            "role": "Backend Engineering Specialist",
            "required_keywords": ["python", "fastapi", "django", "sql", "postgresql", "docker", "rest api", "c++", "java"],
            "suitable_internships": ["Backend Developer Intern", "API Engineering Intern"],
            "summary": "Developing scalable servers, database schemas, and microservices."
        },
        {
            "role": "DevOps & Cloud Infrastructure Engineer",
            "required_keywords": ["docker", "kubernetes", "aws", "linux", "git", "ci/cd", "bash", "python"],
            "suitable_internships": ["Cloud Engineering Intern", "DevOps Intern", "Site Reliability Intern"],
            "summary": "Automating cloud deployments, containers, and server reliability."
        },
        {
            "role": "Cybersecurity & Application Security Analyst",
            "required_keywords": ["linux", "networking", "python", "sql", "cybersecurity", "bash", "vulnerability"],
            "suitable_internships": ["Cybersecurity Intern", "Application Security Intern"],
            "summary": "Assessing software security, vulnerability analysis, and network defense."
        }
    ]
    
    recommendations = []
    for item in roles_catalog:
        matches = [kw for kw in item["required_keywords"] if any(kw in s for s in skills_lower)]
        score = min(98, 60 + len(matches) * 8) if matches else 50
        recommendations.append({
            "role_title": item["role"],
            "match_score": score,
            "matched_skills": [m.capitalize() for m in matches],
            "suitable_internships": item["suitable_internships"],
            "description": item["summary"]
        })
    
    recommendations.sort(key=lambda x: x["match_score"], reverse=True)
    top_tech_skills = all_skills if all_skills else ["Python", "SQL", "Git", "Problem Solving"]
    
    return {
        "candidate_name": resume_data.get("full_name") or "Candidate",
        "top_technical_skills": top_tech_skills[:7],
        "recommended_roles": recommendations[:4],
        "internship_suggestions": recommendations[0]["suitable_internships"] if recommendations else ["Software Development Intern"]
    }

def generate_interview_prep(role_title: str, resume_data: dict, company_name: str = "Target Tech Company") -> dict:
    skills = resume_data.get("skills", []) + resume_data.get("technical_skills", [])
    skills_str = ", ".join(skills) if skills else "Python, SQL, System Design"
    
    tech_questions = [
        {
            "question": f"Explain how you would architect a solution for key requirements in {role_title} using {skills_str[:30]}?",
            "category": "Architecture & Design",
            "answer_guidance": f"Outline system components, data flows, database schemas, and discuss how you leverage {skills_str[:25]} to optimize performance and security."
        },
        {
            "question": f"What are the best practices for error handling, state management, and API reliability when working as a {role_title}?",
            "category": "Core Technical Skills",
            "answer_guidance": "Demonstrate understanding of input validation, HTTP status codes, exception logging, retry logic, and clean code principles."
        },
        {
            "question": "Walk us through a complex bug or performance bottleneck you encountered in one of your projects and how you resolved it.",
            "category": "Problem Solving",
            "answer_guidance": "Use the STAR method (Situation, Task, Action, Result). Highlight diagnostic tooling, root-cause analysis, and measurable performance improvements."
        }
    ]
    
    hr_questions = [
        {
            "question": f"Why are you passionate about pursuing a {role_title} position at {company_name}?",
            "category": "Motivation & Alignment",
            "answer_guidance": f"Connect your personal learning goals and technical project achievements with {company_name}'s mission and engineering values."
        },
        {
            "question": "Describe a scenario where project requirements changed rapidly. How did you adapt your timeline and priorities?",
            "category": "Agility & Communication",
            "answer_guidance": "Focus on clear communication, breaking tasks into sprint deliverables, and managing stakeholder expectations smoothly."
        }
    ]
    
    roadmap = [
        {"phase": "Phase 1: Fundamental Review", "duration": "Days 1-3", "focus": "Brush up on Core CS Fundamentals, Data Structures, Algorithms, and OOP Principles."},
        {"phase": "Phase 2: Role-Specific Deep Dive", "duration": "Days 4-7", "focus": f"Master practical concepts for {role_title}, frameworks ({skills_str[:20]}), and hands-on coding exercises."},
        {"phase": "Phase 3: System Design & Scenario Questions", "duration": "Days 8-10", "focus": "Practice architecture diagrams, API contract design, database schema modeling, and scaling patterns."},
        {"phase": "Phase 4: Behavioral & Mock Interviews", "duration": "Days 11-14", "focus": "Prepare STAR methodology responses for project achievements, team collaboration, and conduct timed mock interviews."}
    ]
    
    topics_to_prepare = [
        f"{role_title} Core Architecture & Frameworks",
        "Data Structures, Algorithms & Time Complexity",
        "Database Design, Indexing & Query Optimization",
        "RESTful API & System Integration Best Practices",
        "Git Workflow, CI/CD & Deployment Basics"
    ]
    
    learning_path = [
        {"topic": "Hands-on Project Building", "recommendation": "Extend one of your portfolio projects with comprehensive tests, containerization (Docker), and clean API documentation."},
        {"topic": "Coding Practice", "recommendation": "Solve 2-3 targeted problem-solving questions daily on array manipulation, hash maps, and dynamic programming."},
        {"topic": "Mock Interviews", "recommendation": "Record yourself explaining technical concepts out loud to refine your verbal technical communication."}
    ]
    
    return {
        "target_role": role_title,
        "company_name": company_name,
        "technical_questions": tech_questions,
        "hr_questions": hr_questions,
        "answer_guidance_summary": "Structure all technical responses with clear assumptions, trade-offs, step-by-step implementation details, and STAR framework stories for behavioral queries.",
        "preparation_roadmap": roadmap,
        "topics_to_prepare": topics_to_prepare,
        "learning_path": learning_path
    }

def generate_doc_qa(doc_text: str, user_question: str = None) -> dict:
    snippets = [line.strip() for line in doc_text.splitlines() if line.strip()]
    summary = " ".join(snippets[:10])[:500] if snippets else "No text extracted from document."
    
    generated_qa = []
    if len(snippets) >= 2:
        generated_qa.append({
            "question": "What is the primary topic or document summary?",
            "answer": f"The document primarily covers: {summary[:250]}..."
        })
    if len(snippets) >= 5:
        generated_qa.append({
            "question": "What are the key technical or structural points mentioned?",
            "answer": f"Key excerpts include: {snippets[min(len(snippets)-1, 3)]} and {snippets[min(len(snippets)-1, 5)]}."
        })
        
    answer_for_user = ""
    if user_question:
        q_lower = user_question.lower()
        matched = [s for s in snippets if any(w in s.lower() for w in q_lower.split() if len(w) > 3)]
        if matched:
            answer_for_user = f"Based on the document context:\n\n" + "\n".join(f"- {m}" for m in matched[:4])
        else:
            answer_for_user = f"Based on the document context, here is a relevant excerpt:\n\n\"{summary[:350]}...\""
            
    return {
        "document_summary": summary,
        "extracted_lines_count": len(snippets),
        "user_question": user_question,
        "user_answer": answer_for_user,
        "generated_qa": generated_qa
    }

def chat_mock_interview(role_title: str, user_message: str, chat_history: list = None, resume_data: dict = None, doc_context: str = None) -> dict:
    if settings.openai_api_key:
        try:
            from openai import OpenAI
            client = OpenAI(api_key=settings.openai_api_key)
            system_prompt = f"You are an intelligent AI Career & Interview Preparation Agent. Memorize all candidate information, resume details ({json.dumps(resume_data or {})}), document context ({doc_context[:1000] if doc_context else 'None'}), and previous chat history. Answer candidate questions accurately using full conversation memory. If the user asks questions about suitable roles, internships, top technical skills, interview preparation, or document context, provide thorough answers. If off-topic, focus on career and interview guidance."
            messages = [{"role": "system", "content": system_prompt}]
            if chat_history:
                for h in chat_history:
                    messages.append({"role": "user" if (h.get("sender") == "user" or h.get("sender") == "candidate") else "assistant", "content": h.get("text", "")})
            messages.append({"role": "user", "content": user_message})
            response = client.chat.completions.create(
                model=settings.openai_model,
                temperature=0.5,
                messages=messages
            )
            return {
                "response": response.choices[0].message.content,
                "role_title": role_title
            }
        except Exception as e:
            pass
            
    # Smart local Preparation Agent
    user_msg_lower = user_message.lower().strip()
    
    # Extract candidate user messages from history
    user_messages = [
        h.get("text") for h in (chat_history or []) 
        if (h.get("sender") == "user" or h.get("sender") == "candidate") and h.get("text") and h.get("text").strip() != user_message
    ]
    
    words = set(re.findall(r"\b[a-z0-9]+\b", user_msg_lower))
    cand_name = resume_data.get("full_name") or resume_data.get("name") or "Candidate"
    skills = resume_data.get("skills", []) + resume_data.get("technical_skills", [])
    if not skills:
        skills = ["Python", "FastAPI", "SQL", "React", "Docker", "Git"]
    skills_unique = list(dict.fromkeys(skills))

    # 1. Document-Based Q&A check if doc_context is present
    if doc_context and any(kw in user_msg_lower for kw in ["document", "pdf", "docx", "file", "excerpt", "summary", "text", "context"]):
        doc_lines = [l.strip() for l in doc_context.splitlines() if l.strip()]
        matched = [l for l in doc_lines if any(w in l.lower() for w in user_msg_lower.split() if len(w) > 3)]
        if matched:
            reply = f"Based on the uploaded document context:\n" + "\n".join(f"• {m}" for m in matched[:4])
        else:
            snippet = " ".join(doc_lines[:5])[:300]
            reply = f"Here is the relevant excerpt from your uploaded document:\n\"{snippet}...\""

    # 2. Specific Resume-Based Questions: "Which role can I apply for?" / "What role fits me?"
    elif any(phrase in user_msg_lower for phrase in ["which role", "what role", "roles can i apply", "suitable role", "roles for me", "recommend role"]):
        rec_data = recommend_roles(resume_data or {})
        top_roles = rec_data.get("recommended_roles", [])
        roles_summary = "\n".join(f"{idx+1}. **{r['role_title']}** ({r['match_score']}% match) — Key skills matched: {', '.join(r['matched_skills'][:4])}" for idx, r in enumerate(top_roles[:3]))
        reply = f"Hello {cand_name}! Based on your extracted resume skills ({', '.join(skills_unique[:5])}), here are the top roles you are best suited for:\n\n{roles_summary}\n\nYou can select any of these roles to generate a complete interview prep plan!"

    # 3. "Which internship is suitable for my skills?"
    elif any(phrase in user_msg_lower for phrase in ["which internship", "internship is suitable", "suitable internship", "internships for my skills", "best internship"]):
        rec_data = recommend_roles(resume_data or {})
        top_role = rec_data.get("recommended_roles", [{}])[0]
        internships = top_role.get("suitable_internships", ["AI/ML Intern", "Software Engineering Intern", "Backend Developer Intern"])
        reply = f"Based on your profile ({cand_name}, skills in {', '.join(skills_unique[:4])}), the most suitable internship opportunities for you are:\n\n" + "\n".join(f"• **{item}**" for item in internships) + f"\n\nThese align closely with your top matched role **{top_role.get('role_title', 'Software Engineer')}**!"

    # 4. "What are my strongest technical skills?"
    elif any(phrase in user_msg_lower for phrase in ["strongest", "technical skill", "my skills", "top skills", "key skills"]):
        reply = f"Based on analysis of your extracted resume data, your top technical skills are:\n\n" + "\n".join(f"• **{s}**" for s in skills_unique[:7]) + f"\n\nThese form a strong foundation for technical roles like {role_title or 'Software Developer'}!"

    # 5. Session Memory / First Question Queries
    elif any(phrase in user_msg_lower for phrase in ["what was my first question", "what did i ask first", "first question", "remember my first"]):
        if user_messages and user_messages[0]:
            reply = f"Your first question in this session was: '{user_messages[0].strip()}'"
        else:
            reply = f"Your first question was: '{user_message}'"

    # 6. Interview Prep / Roadmap Queries
    elif any(kw in user_msg_lower for kw in ["interview", "prep", "roadmap", "question", "technical", "hr"]):
        reply = f"For **{role_title}**, you should prepare key areas:\n1. Core CS & System Design using {', '.join(skills_unique[:3])}.\n2. Behavioral STAR stories detailing your top projects.\n3. Interactive mock questions. Check the structured cards on this screen for questions and answer guidance!"

    # 7. Greetings
    elif any(g in words for g in ["hello", "hi", "hii", "hey", "start", "greetings"]):
        reply = "Hello! How can I help you with your career goals and interview preparation today?"

    # 8. Unwanted / Off-Topic Question Handling
    else:
        reply = "I am your Career & Internship Assistant! Please ask questions related to internships, job roles, technical skills, resume analysis, or interview preparation."

    return {
        "response": reply,
        "role_title": role_title
    }


